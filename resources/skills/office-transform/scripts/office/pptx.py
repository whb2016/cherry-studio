"""Everything pptx: the extract path. There is no pptx edit path — patch-copy handles xlsx and
docx only.

python-pptx is imported inside the functions, never at module scope: office_extract.py imports
this module eagerly for every format, and a top-level third-party import would make a
pptx-less environment fail on an xlsx source. The in-function `from pptx import ...` is an
absolute import and resolves to the third-party package, not to this module.
"""

from pathlib import Path

from office.common import fail, require_index


def iter_shapes_recursive(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes_recursive(shape.shapes)


def shape_text_lines(shape) -> list[str]:
    if shape.has_text_frame:
        return [paragraph.text for paragraph in shape.text_frame.paragraphs]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    return []


def extract_pptx(src: Path, anchor: dict, out_path: Path, out_format: str) -> None:
    try:
        from pptx import Presentation
    except ImportError:
        fail("python-pptx is required for pptx sources — rerun via `uv run --with python-pptx python ...`")

    if anchor.get("slide") is None:
        fail("pptx anchor requires a one-based 'slide' number")
    slide_number = require_index(anchor.get("slide"), "pptx anchor 'slide'", 1)
    slide_index = slide_number - 1

    presentation = Presentation(str(src))
    slides = list(presentation.slides)
    if slide_index >= len(slides):
        fail(f"slide {slide_number} out of range (deck has {len(slides)} slides)")
    slide = slides[slide_index]

    node_id = anchor.get("nodeId")
    if node_id is None:
        if anchor.get("paragraph") is not None or anchor.get("tableCell") is not None:
            fail("pptx anchor has 'paragraph'/'tableCell' but no 'nodeId'; refusing to fall back to whole-slide extraction")
        lines = [line for shape in iter_shapes_recursive(slide.shapes) for line in shape_text_lines(shape)]
    else:
        shape = next(
            (candidate for candidate in iter_shapes_recursive(slide.shapes) if str(candidate.shape_id) == str(node_id)),
            None,
        )
        if shape is None:
            fail(f"shape with nodeId {node_id!r} not found on slide {slide_number}")

        table_cell = anchor.get("tableCell")
        paragraph_index = anchor.get("paragraph")
        if table_cell is not None and paragraph_index is not None:
            fail("pptx anchor has both 'paragraph' and 'tableCell'; they address different things — pick one")
        if table_cell is not None:
            if not isinstance(table_cell, dict):
                fail(f"pptx anchor 'tableCell' must be an object with 'row' and 'col': {table_cell!r}")
            if not (getattr(shape, "has_table", False) and shape.has_table):
                fail(f"shape {node_id!r} is not a table but anchor has 'tableCell'")
            rows = list(shape.table.rows)
            row = require_index(table_cell.get("row"), "pptx anchor tableCell 'row'", 0)
            col = require_index(table_cell.get("col"), "pptx anchor tableCell 'col'", 0)
            if row >= len(rows) or col >= len(list(rows[row].cells)):
                fail(f"tableCell {table_cell!r} out of range for shape {node_id!r}")
            lines = [list(rows[row].cells)[col].text]
        elif paragraph_index is not None:
            if not shape.has_text_frame:
                fail(f"shape {node_id!r} has no text body but anchor has 'paragraph'")
            paragraphs = shape.text_frame.paragraphs
            # Bounded on both sides: a negative ordinal would index from the end and quietly return a
            # paragraph nobody asked for.
            paragraph_index = require_index(paragraph_index, "pptx anchor 'paragraph'", 0)
            if paragraph_index >= len(paragraphs):
                fail(f"paragraph {paragraph_index} out of range (shape has {len(paragraphs)} paragraphs)")
            lines = [paragraphs[paragraph_index].text]
        else:
            lines = shape_text_lines(shape)

    if out_format in ("txt", "md"):
        out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    else:
        fail(f"unsupported output format for pptx source: {out_format!r} (use txt or md)")
