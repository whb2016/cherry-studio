"""Everything xlsx: the A1 primitives and the extract path.

The user-visible description of what an xlsx edit does lives in ../office_patch_copy.py's
module docstring — that text is the script's --help output, so changing behaviour here means
changing it there too.

openpyxl is imported inside extract_xlsx, never at module scope: office_extract.py imports
this module eagerly for every format, and a top-level third-party import would make an
xlsx-only environment fail on a docx source (and vice versa).
"""

import csv
import datetime
import re
from pathlib import Path

from office.common import fail

A1_CELL_RE = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")

MAX_RANGE_CELLS = 1_000_000

# The SpreadsheetML grid (ECMA-376): columns A..XFD, rows 1..1048576.
MAX_COLUMN_INDEX = 16_384
MAX_ROW_NUMBER = 1_048_576


def column_to_index(letters: str) -> int:
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def parse_a1_cell(ref: str) -> tuple[int, int]:
    match = A1_CELL_RE.match(ref)
    if not match:
        fail(f"invalid A1 cell reference: {ref!r}")
    column, row = column_to_index(match.group(1)), int(match.group(2))
    if column > MAX_COLUMN_INDEX or row > MAX_ROW_NUMBER:
        fail(f"cell {ref!r} is outside the worksheet grid (max XFD{MAX_ROW_NUMBER})")
    return column, row


def parse_a1_range(ref: str) -> tuple[int, int, int, int]:
    """Return (min_col, min_row, max_col, max_row) from 'B2' or 'A1:C10'."""
    parts = ref.split(":")
    if len(parts) > 2:
        fail(f"invalid A1 range: {ref!r}")
    start = parse_a1_cell(parts[0])
    end = parse_a1_cell(parts[-1])
    return (
        min(start[0], end[0]),
        min(start[1], end[1]),
        max(start[0], end[0]),
        max(start[1], end[1]),
    )


def cell_display(value) -> str:
    """Render a cell in a shape a spreadsheet reader recognises, not the way Python prints it.

    csv and md output is read as spreadsheet text and compared against the renderer's excerpt, so a
    `str()` form only Python uses is wrong on both counts: `True` where every spreadsheet writes
    `TRUE`, `2024-01-03 00:00:00` for a cell the user sees as `2024-01-03` (openpyxl hands back a
    datetime for date-only cells too, never a bare date), and `1 day, 2:30:00` for a duration Excel
    counts the hours through as `26:30:00`. Numbers keep their stored value (`0.4567`, not `45.67%`):
    presenting them any other way means implementing number formats, which SKILL.md documents as an
    accepted asymmetry between the extract and the excerpt.

    A midnight time is what marks a date-only cell, because the cell's number format is not carried
    through extract_xlsx and threading it there to decide this would touch every output path. The
    rule errs in both directions — a date-formatted cell that stores a time prints the time, a
    date-time cell that stores midnight prints as a date — and both are the number-format difference
    the anchor check already sets aside. A text cell that says `TRUE` renders the same as the
    boolean; so does the renderer, which is the point.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, datetime.time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, datetime.timedelta):
        # Floor division on a negative total borrows an hour (-30 minutes would print -1:30:00), so
        # split the sign off first and format the magnitude.
        seconds = int(value.total_seconds())
        sign = "-" if seconds < 0 else ""
        seconds = abs(seconds)
        return f"{sign}{seconds // 3600}:{seconds // 60 % 60:02d}:{seconds % 60:02d}"
    return str(value)


def write_markdown_table(rows: list[list[str]], out_path: Path) -> None:
    if not rows:
        fail("selection produced no rows")
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    # Backslashes go first. Escaping only the pipe turns a cell's own `\` into the escape for the
    # pipe that follows it, so `a\|b` reaches the reader as an escaped backslash and a live
    # separator — one cell silently becomes two.
    escaped = [
        [cell.replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ") for cell in row] for row in normalized
    ]
    lines = ["| " + " | ".join(escaped[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in escaped[1:])
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def extract_xlsx(src: Path, anchor: dict, out_path: Path, out_format: str) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        fail("openpyxl is required for xlsx sources — rerun via `uv run --with openpyxl python ...`")

    sheet_name = anchor.get("sheet")
    range_ref = anchor.get("range")
    if not sheet_name or not range_ref:
        fail("xlsx anchor requires 'sheet' and 'range'")
    # 'range' is one A1 string. The two-element form charRange uses in the same anchor would reach
    # .split() and traceback, and writing the pair here rather than the string is an easy slip.
    if not isinstance(range_ref, str):
        fail(f"xlsx anchor 'range' must be an A1 string like 'A1:C10', not {type(range_ref).__name__}: {range_ref!r}")

    min_col, min_row, max_col, max_row = parse_a1_range(range_ref)
    area = (max_row - min_row + 1) * (max_col - min_col + 1)
    if area > MAX_RANGE_CELLS:
        fail(f"range {range_ref!r} covers {area} cells (limit {MAX_RANGE_CELLS}); select a smaller region")

    workbook = load_workbook(src, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        fail(f"worksheet not found: {sheet_name!r} (has: {workbook.sheetnames})")
    worksheet = workbook[sheet_name]

    values = []
    # Derived-sheet coordinates of the cells that really hold an error value (`#N/A`). Under data_only
    # an error and the text "#N/A" read back as the same string; only the cell's data_type still says
    # which it was, and the xlsx output below needs to know.
    error_cells = set()
    rows = worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
    for row_offset, row in enumerate(rows):
        values.append([cell.value for cell in row])
        error_cells.update(
            (row_offset + 1, column_offset + 1)
            for column_offset, cell in enumerate(row)
            if cell.data_type == "e"
        )
    # Merge followers are not masked: a read_only worksheet has no merged_cells, so a mask keyed on it
    # never runs, and fetching the ranges means either a second, non-streaming load — giving up the
    # streaming this reader exists for — or hand-parsing <mergeCells>. Excel and openpyxl clear a
    # follower when the merge is made, so followers read back empty and match what the user sees;
    # the file that kept hidden text under a merge extracts it, which SKILL.md "## Limits" says out
    # loud. Any future mask must clamp to the rows iter_rows actually returned, not to max_row:
    # read_only stops at the last populated row, so a merge below the data would index past `values`.
    workbook.close()

    if out_format == "xlsx":
        derived = Workbook()
        derived_sheet = derived.active
        derived_sheet.title = sheet_name[:31]
        for row in values:
            derived_sheet.append(row)
        # append() re-infers each cell's type from its value: a string starting with "=" becomes a
        # formula and one of Excel's error codes (`#N/A`) becomes an error, so text the source merely
        # displayed comes back as something the spreadsheet runs or reports — and the formula has no
        # cached value, so re-extracting it reads nothing. Put such a cell back to a string, with
        # quotePrefix so Excel keeps treating it as text after someone edits it. A cell that held a
        # real error keeps it; error_cells is what tells the two apart.
        for row in derived_sheet.iter_rows():
            for cell in row:
                if cell.data_type not in ("f", "e"):
                    continue
                if cell.data_type == "e" and (cell.row, cell.column) in error_cells:
                    continue
                cell.data_type = "s"
                cell.quotePrefix = True
        derived.save(out_path)
    elif out_format == "csv":
        with out_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows([[cell_display(value) for value in row] for row in values])
    elif out_format == "md":
        write_markdown_table([[cell_display(value) for value in row] for row in values], out_path)
    else:
        fail(f"unsupported output format for xlsx source: {out_format!r} (use xlsx, csv, or md)")
